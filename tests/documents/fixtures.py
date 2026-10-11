"""Synthetic test documents (no real data): every value is a fixture literal shared with the M0 tests."""

from __future__ import annotations

import dataclasses
import io

from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import FreeText
from pypdf.generic import RectangleObject
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from dps.layers.documents.render import PdfRenderer

KYC_LINES = (
    "Customer: Ananya Krishnan",
    "SSN 123-45-6789 on file",
    "Email ananya.k@example-mail.in",
    "Card 4111 1111 1111 1111 expires 09/28",
)


def text_pdf(
    lines=KYC_LINES, *, pages: int = 1, size=(612.0, 792.0), font_size: int = 14, x: float = 72.0
) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=size, invariant=1)
    for n in range(pages):
        c.setFont("Helvetica", font_size)
        for i, line in enumerate(lines):
            c.drawString(
                x, size[1] - 72 - i * font_size * 2, line if n == 0 else f"Page {n + 1}: {line}"
            )
        c.showPage()
    c.save()
    return buf.getvalue()


def rewrite(
    data: bytes, *, rotate: int = 0, crop=None, encrypt=None, author=None, annotation=None
) -> bytes:
    reader = PdfReader(io.BytesIO(data))
    writer = PdfWriter()
    for page in reader.pages:
        if rotate:
            page.rotation = rotate
        if crop is not None:
            page.cropbox = RectangleObject(crop)
        writer.add_page(page)
    if author is not None:
        writer.add_metadata({"/Author": author})
    if annotation is not None:
        writer.add_annotation(
            page_number=0, annotation=FreeText(text=annotation, rect=(50, 550, 300, 600))
        )
    if encrypt is not None:
        writer.encrypt(user_password=encrypt, owner_password="owner")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def rotated_pdf(lines=KYC_LINES) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(612, 792), invariant=1)
    c.saveState()
    c.translate(612, 0)
    c.rotate(90)
    c.setFont("Helvetica", 14)
    for i, line in enumerate(lines):
        c.drawString(72, 500 - i * 28, line)
    c.restoreState()
    c.showPage()
    c.save()
    return rewrite(buf.getvalue(), rotate=90)


def cropped_pdf(lines=KYC_LINES) -> bytes:
    return rewrite(text_pdf(lines, x=150), crop=(100, 100, 500, 700))


def encrypted_pdf(password: str = "pw") -> bytes:
    return rewrite(text_pdf(), encrypt=password)


def annotated_pdf() -> bytes:
    return rewrite(text_pdf(), author="Ananya Krishnan", annotation="SSN 123-45-6789")


def scanned_pdf(lines=KYC_LINES, *, dpi: int = 200, pages: int = 1) -> bytes:
    src = text_pdf(lines, pages=pages)
    buf = io.BytesIO()
    c = canvas.Canvas(buf, invariant=1)
    with PdfRenderer(src) as renderer:
        for i in range(len(renderer)):
            width, height = renderer.page_size(i)
            c.setPageSize((width, height))
            c.drawImage(ImageReader(renderer.render(i, dpi)), 0, 0, width=width, height=height)
            c.showPage()
    c.save()
    return buf.getvalue()


def hybrid_pdf(lines=KYC_LINES) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(612, 792), invariant=1)
    c.setFont("Helvetica", 14)
    for i, line in enumerate(lines):
        c.drawString(72, 720 - i * 28, line)
    c.drawImage(ImageReader(Image.new("RGB", (800, 600), "white")), 72, 72, width=468, height=351)
    c.showPage()
    c.save()
    return buf.getvalue()


def small_image_pdf(lines=KYC_LINES) -> bytes:
    """A text page with one small image (100 x 60 pt, under 2 % of the page): too little image
    for ``ocr_image_coverage`` to make the page scan-suspect."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(612, 792), invariant=1)
    c.setFont("Helvetica", 14)
    for i, line in enumerate(lines):
        c.drawString(72, 720 - i * 28, line)
    c.drawImage(ImageReader(Image.new("RGB", (200, 120), "white")), 400, 100, width=100, height=60)
    c.showPage()
    c.save()
    return buf.getvalue()


#: An invisible text layer with OCR errors, as legacy scanners write under a page image.
GARBLED_LAYER = (
    "Customer: Anonymous Person",
    "SSN 1Z3-4S-67B9 on file",
    "Email ananya.k at example-mail",
    "Card 4lll llll llll llll expires O9/2B",
)
#: An invisible text layer that says something else entirely at the start of every line.
SPECIMEN_LAYER = ("Specimen",) * 4


def searchable_scan(hidden=GARBLED_LAYER, lines=KYC_LINES, *, dpi: int = 200) -> bytes:
    """A full-page image of ``lines`` under an invisible (render mode 3) text layer that says
    ``hidden`` at the same line positions: a searchable scan whose text layer is not what it
    shows."""
    with PdfRenderer(text_pdf(lines)) as renderer:
        image = renderer.render(0, dpi)
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(612, 792), invariant=1)
    c.drawImage(ImageReader(image), 0, 0, width=612, height=792)
    text = c.beginText()
    text.setTextRenderMode(3)
    text.setFont("Helvetica", 14)
    for i, line in enumerate(hidden):
        text.setTextOrigin(72, 792 - 72 - i * 28)
        text.textLine(line)
    c.drawText(text)
    c.showPage()
    c.save()
    return buf.getvalue()


def broken_tounicode_pdf(lines=KYC_LINES) -> bytes:
    """``text_pdf(lines)`` whose fonts carry a ``/ToUnicode`` map that reads every digit as ``X``:
    the page shows the digits, its text layer does not (as in many generated statements)."""
    from pypdf.generic import NameObject, StreamObject

    entries = [
        f"<{code:02X}> <{(0x58 if chr(code).isdigit() else code):04X}>"
        for code in range(0x20, 0x7F)
    ]
    cmap = StreamObject()
    cmap.set_data(
        "\n".join(
            [
                "/CIDInit /ProcSet findresource begin",
                "12 dict begin",
                "begincmap",
                "/CMapName /BrokenDigits def",
                "1 begincodespacerange <00> <FF> endcodespacerange",
                f"{len(entries)} beginbfchar",
                *entries,
                "endbfchar",
                "endcmap",
                "CMapName currentdict /CMap defineresource pop",
                "end",
                "end",
            ]
        ).encode()
    )
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(text_pdf(lines))))
    ref = writer._add_object(cmap)
    for page in writer.pages:
        fonts = page["/Resources"]["/Font"]
        for name in fonts:
            fonts[name].get_object()[NameObject("/ToUnicode")] = ref
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def png_document(lines=KYC_LINES, *, dpi: int = 200) -> bytes:
    with PdfRenderer(text_pdf(lines)) as renderer:
        image = renderer.render(0, dpi)
    buf = io.BytesIO()
    image.save(buf, "PNG", dpi=(dpi, dpi))
    return buf.getvalue()


def jpeg_rotated_exif(width: int = 300, height: int = 100) -> bytes:
    exif = Image.Exif()
    exif[0x0112] = 6  # orientation: rotate 90 CW on display
    buf = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(buf, "JPEG", exif=exif.tobytes())
    return buf.getvalue()


def huge_png(width: int = 8000, height: int = 6000) -> bytes:
    buf = io.BytesIO()
    Image.new("1", (width, height), 1).save(buf, "PNG")
    return buf.getvalue()


def docx_document() -> bytes:
    import docx

    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Ref CUST-884120"
    document.add_paragraph(KYC_LINES[0])
    table = document.add_table(rows=2, cols=2)
    for r, (label, value) in enumerate((("SSN", "123-45-6789"), ("Card", "4111 1111 1111 1111"))):
        table.cell(r, 0).text = label
        table.cell(r, 1).text = value
    paragraph = document.add_paragraph("Mail ananya.k@example-mail.in")
    paragraph.add_run().add_break()
    paragraph.add_run("call +91 98451 20377")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def csv_document(rows: list[dict[str, str]], *, delimiter: str = ",") -> bytes:
    import csv

    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(rows[0]), delimiter=delimiter, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


CUSTOMERS = [
    {
        "customer_id": "CUST-884120",
        "full_name": "Ananya Krishnan",
        "email": "ananya.k@example-mail.in",
        "notes": "card 4111 1111 1111 1111, call later",
        "risk_score": "712",
    },
    {
        "customer_id": "CUST-884121",
        "full_name": "Rohan Desai",
        "email": "rohan.d@example-mail.in",
        "notes": "no remarks",
        "risk_score": "640",
    },
]


class FakeOcr:
    """An ``OcrClient`` double: each ``read`` answers with the next prepared page of words.

    Words are in pixels of the image DPS sends (``pdf_render_dpi`` for PDF pages); the page size
    is that image's, as Read v3.2 reports it."""

    engine = "fake-ocr"

    def __init__(
        self,
        pages: list[list[tuple[str, tuple[float, float, float, float]]]],
        *,
        confidence: float = 0.97,
    ) -> None:
        self.pages = pages
        self.calls = 0
        self.images: list[bytes] = []
        self.confidence = confidence

    def configured(self) -> bool:
        return True

    async def read(self, image: bytes, *, content_type: str = "image/png"):  # noqa: ANN201
        from dps.layers.documents.ocr import OcrPage, OcrWord

        assert image.startswith(b"\x89PNG\r\n\x1a\n")
        self.images.append(image)
        words = self.pages[min(self.calls, len(self.pages) - 1)] if self.pages else []
        self.calls += 1
        with Image.open(io.BytesIO(image)) as sent:
            width, height = sent.size
        return OcrPage(
            width=float(width),
            height=float(height),
            unit="pixel",
            angle=0.0,
            words=[
                OcrWord(
                    text=t,
                    polygon=(x0, y0, x1, y0, x1, y1, x0, y1),
                    confidence=self.confidence,
                    line=int(y0 // 20),
                )
                for t, (x0, y0, x1, y1) in words
            ],
        )

    async def probe(self) -> tuple[bool, str]:
        return True, "up"

    async def aclose(self) -> None:
        return None


class ScreenOcr(FakeOcr):
    """An OCR double that reads what an image *shows*: the words drawn on the page (pixels at the
    size of the image sent; the same words for every page), less every word a redaction box hides
    (its area is all but black in the image) — or, with ``box_reads_as``, that text in its place,
    as a real OCR sometimes reads a black box as letters. ``reads`` keeps the texts each call
    returned."""

    def __init__(
        self,
        drawn: list[tuple[str, tuple[float, float, float, float]]],
        *,
        box_reads_as: str | None = None,
        confidence: float = 0.97,
    ) -> None:
        super().__init__([drawn], confidence=confidence)
        self.box_reads_as = box_reads_as
        self.reads: list[list[str]] = []

    async def read(self, image: bytes, *, content_type: str = "image/png"):  # noqa: ANN201
        import math

        from PIL import ImageStat

        from dps.layers.documents.ocr import OcrPage

        page = await super().read(image, content_type=content_type)
        with Image.open(io.BytesIO(image)) as sent:
            grey = sent.convert("L")
        shown = []
        for word in page.words:
            xs, ys = word.polygon[0::2], word.polygon[1::2]
            region = grey.crop(
                (math.floor(min(xs)), math.floor(min(ys)), math.ceil(max(xs)), math.ceil(max(ys)))
            )
            if ImageStat.Stat(region).mean[0] >= 64:  # painted boxes are black
                shown.append(word)
            elif self.box_reads_as is not None:
                shown.append(dataclasses.replace(word, text=self.box_reads_as))
        self.reads.append([w.text for w in shown])
        return OcrPage(
            width=page.width, height=page.height, unit=page.unit, angle=page.angle, words=shown
        )


def words_px(
    pdf: bytes, *, dpi: int = 200, drop: tuple[str, ...] = ()
) -> list[list[tuple[str, tuple[float, float, float, float]]]]:
    """Text-layer words of ``pdf`` per page in pixels at ``dpi``: what a perfect OCR returns for the
    scanned twin of ``pdf``. ``drop`` removes words (e.g. those a redaction painted over)."""
    from dps.config import Settings
    from dps.layers.documents.extract import extract_pdf

    doc = extract_pdf(pdf, Settings(_env_file=None), "auto")
    s = dpi / 72.0
    return [
        [
            (w.text, (w.bbox.x0 * s, w.bbox.y0 * s, w.bbox.x1 * s, w.bbox.y1 * s))
            for w in p.words
            if w.text not in drop
        ]
        for p in doc.pages
    ]

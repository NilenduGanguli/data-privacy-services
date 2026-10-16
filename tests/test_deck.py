"""The showcase deck (M5 Task 26B): built from the demo's value-free outputs, Keynote-safe.

Every figure on a slide must come from an input (the e2e demo report, the bench JSON); a missing
input renders a slate placeholder box saying what is missing, never a made-up number. Keynote (the
owner's viewer) drops PowerPoint chart parts and lacks some fonts, so the deck draws charts as plain
shapes and names only Helvetica Neue and Arial.
"""

from __future__ import annotations

import json
import re
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest

pytest.importorskip("pptx")

from PIL import Image  # noqa: E402  (pptx imported above implies Pillow)
from pptx import Presentation  # noqa: E402

from dps.engine.detect.ensemble import Detector, DetectorConfig  # noqa: E402
from dps.engine.taxonomy import Sensitivity  # noqa: E402
from tools import build_deck  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "docs" / "demo" / "DEMO-SCRIPT.md"
ALLOWED_FONTS = {"Helvetica Neue", "Arial"}
#: A figure: digits not glued to a letter, digit, underscore or dot (labels such as S8, L2, D4, P1,
#: e2e and v1 are names, not figures).
NUMBER = re.compile(r"(?<![A-Za-z0-9_.])\d[\d,]*(?:\.\d+)?")

BENCH = {
    "batch_rows": 1000,
    "engine": "oracle",
    "engine_version": "0.2.0",
    "fpe_enabled": True,
    "host": {"cpus": 18, "machine": "arm64"},
    "job_id": "j_01M4VSVE8QSM41YZT9B9N83PH8",
    "label": "1-worker",
    "mode": "obfuscate",
    "parallelism": 4,
    "partition_rows_per_second": {"p50": 246.9, "p95": 246.9},
    "partitions": 4,
    "rows": 20000,
    "rows_per_second": 941.6,
    "state": "succeeded",
    "wall_seconds": 21.24,
    "workers_online": 1,
    "write_mode": "table",
}

S8_FACTS: dict[str, Any] = {
    "counts": {"rows_written": 1987},
    "fidelity": {"score": 0.9312},
    "privacy": {
        "exact_row_matches": 0,
        "nndr": {"ratio": 1.163, "target": 1.0, "measured": 1, "passed": True},
        "membership_inference": {"run": True, "auc": 0.5134, "target": 0.6, "passed": True},
    },
    "utility": {"run": True, "metric": "accuracy", "ratio": 0.9721},
}
#: S9 as the smoke's D4 check measures it (tools/smoke_test.py, Task 24 S4b): the scan's evidence
#: block, how far the demo site's own /__received counters moved across the scan, and the sends the
#: request guard stopped (the two positive controls: the site served the scan, its pages tried to
#: send).
S9_FACTS: dict[str, Any] = {
    "scan": {"pages": 21, "fields": 63, "personas": 3, "submissions": 0, "non_get_allowed": 0},
    "received": {"non_get": 0, "form_get": 0, "logout_hits": 0, "admin_hits": 0, "get": 175},
    "sends_blocked": 60,
}
FULL_FACTS: dict[str, dict[str, Any]] = {
    "S1": {"raw_pii_sent_to_llm": 0},
    "S2": {
        "findings": 13,
        "entity_types": ["AADHAAR", "EMAIL", "PAN_IN", "PERSON", "PHONE"],
        "reveal_identical": True,
        "raw_pii_sent_to_llm": 0,
    },
    "S3": {},
    "S4": {"ocr_pages": 2, "findings": 9, "extractable_chars": 0, "originals_left": 0},
    "S5": {},
    "S6": {
        "jobs": [
            {
                "engine": "oracle",
                "counts": {"rows_written": 2000, "partitions": 4},
                "timings": {"rows_per_s": 812.4},
                "verification": {"passed": True},
            },
            {
                "engine": "mongo",
                "counts": {"rows_written": 2000, "partitions": 4},
                "timings": {"rows_per_s": 1430.2},
                "verification": {"passed": True},
            },
            {
                "engine": "postgres",
                "counts": {"rows_written": 2000, "partitions": 4},
                "timings": {"rows_per_s": 2210.9},
                "verification": {"passed": True},
            },
        ]
    },
    "S7": {
        "members": [{"engine": "oracle"}, {"engine": "mongo"}, {"engine": "postgres"}],
        "pairs": [
            {"left": 0, "right": 1, "checked": 500, "matched": 500, "consistent": True},
            {"left": 0, "right": 2, "checked": 500, "matched": 500, "consistent": True},
            {"left": 1, "right": 2, "checked": 500, "matched": 500, "consistent": True},
        ],
        "sampled": 500,
        "all_members_matched": 500,
        "passed": True,
    },
    "S8": S8_FACTS,
    "S9": S9_FACTS,
    "S10": {"packs_listed": 5, "audit_rows": 42, "pdf_ok": True},
    "S11": {
        # the local stack, measured: evaluated as if production, enforced only in staging/prod
        "posture": {"passed": 5, "total": 11, "enforced": False, "env": "local"},
        "rotated_key_status": 401,
        "retention_state": "succeeded",
    },
}
SCREENS_ALL = ("playground", "job", "scan-parallel", "evidence")


# -- fixtures -------------------------------------------------------------------------------------
def _report(steps: dict[str, dict[str, Any]], **top: Any) -> dict[str, Any]:
    return {
        "schema": "dps.e2e-demo.v1",
        "generated_at": "2026-10-14T15:04:05Z",
        "dps_version": "1.0.0",
        "commit": "64c46e3",
        "tenant": "retail-banking-apac",
        "total_seconds": 431.6,
        "steps": [
            {"id": sid, "title": f"step {sid}", "status": "PASS", "seconds": 12.5, "facts": facts}
            for sid, facts in steps.items()
        ],
        **top,
    }


def _inputs(
    tmp_path: Path,
    steps: dict[str, dict[str, Any]],
    *,
    screens: tuple[str, ...] = ("playground", "scan-parallel"),
    bench: bool = True,
    report_top: dict[str, Any] | None = None,
) -> dict[str, Path]:
    report = tmp_path / "e2e-demo-report.json"
    report.write_text(json.dumps(_report(steps, **(report_top or {}))), encoding="utf-8")
    shots = tmp_path / "demo-screens"
    shots.mkdir()
    for name in screens:
        Image.new("RGB", (8, 8), (15, 118, 110)).save(shots / f"{name}-light.png")
    bench_dir = tmp_path / "bench"
    bench_dir.mkdir()
    if bench:
        (bench_dir / "m3-2026-10-13-oracle-obfuscate-table-1-worker.json").write_text(
            json.dumps(BENCH), encoding="utf-8"
        )
    return {"report": report, "screens": shots, "bench": bench_dir}


def _build(tmp_path: Path, paths: dict[str, Path]) -> tuple[Any, build_deck.BuildResult]:
    out = tmp_path / "deck.pptx"
    result = build_deck.build_deck(
        out=out,
        report=paths["report"],
        screens=paths["screens"],
        bench=paths["bench"],
        script=SCRIPT,
    )
    return Presentation(str(out)), result


def _shapes(container: Any) -> Iterator[Any]:
    for shape in container.shapes:
        yield shape
        if shape.shape_type is not None and hasattr(shape, "shapes"):  # group
            yield from _shapes(shape)


def _texts(slide: Any, *, chrome: bool = False) -> list[str]:
    return [
        s.text_frame.text
        for s in _shapes(slide)
        if s.has_text_frame and (chrome or not s.name.startswith("chrome:"))
    ]


def _input_numbers(*docs: Any) -> list[float]:
    def walk(obj: Any) -> Iterator[float]:
        if isinstance(obj, bool):
            return
        if isinstance(obj, int | float):
            yield float(obj)
        elif isinstance(obj, dict):
            for v in obj.values():
                yield from walk(v)
        elif isinstance(obj, list):
            for v in obj:
                yield from walk(v)

    return [n for doc in docs for n in walk(doc)]


def _traceable(token: str, numbers: list[float]) -> bool:
    whole, _, frac = token.rstrip(",").replace(",", "").partition(".")
    printed = f"{int(whole)}.{frac}" if frac else str(int(whole))
    return any(f"{v:.{len(frac)}f}" == printed for v in numbers)


def _names(slide: Any) -> set[str]:
    return {s.name for s in _shapes(slide)}


# -- the plan's acceptance (Task 26B step 1) ------------------------------------------------------
def test_deck_from_fixture_inputs(tmp_path: Path) -> None:
    paths = _inputs(tmp_path, {"S8": S8_FACTS, "S9": S9_FACTS})
    prs, result = _build(tmp_path, paths)
    slides = list(prs.slides)

    # 16:9, exactly twelve slides, each with a title placeholder and speaker notes
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01
    assert len(slides) == 12 and result.slides == 12
    for slide in slides:
        assert slide.shapes.title is not None and slide.shapes.title.text.strip()
        assert slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()

    # no chart parts (Keynote drops them), no embedded workbooks, no graphic frames at all
    with zipfile.ZipFile(result.out) as z:
        names = z.namelist()
        assert not [n for n in names if n.startswith(("ppt/charts/", "ppt/embeddings/"))]
        xml = {n: z.read(n).decode("utf-8") for n in names if n.endswith(".xml")}
    for name, text in xml.items():
        if name.startswith("ppt/slides/slide"):
            assert "<p:graphicFrame" not in text, name
        # only Helvetica Neue and Arial are ever named (the theme references +mj-lt and +mn-lt
        # resolve to the theme's fonts, checked below)
        faces = set(re.findall(r'typeface="([^"]*)"', text))
        assert faces <= ALLOWED_FONTS | {""} | {f for f in faces if f.startswith("+m")}, name
    themes = [t for n, t in xml.items() if n.startswith("ppt/theme/")]
    assert themes
    for theme in themes:
        major_minor = re.findall(r"<a:(?:major|minor)Font>\s*<a:latin typeface=\"([^\"]+)\"", theme)
        assert major_minor and set(major_minor) <= ALLOWED_FONTS

    # every figure printed on slides 7-11 comes from an input (no invented numbers)
    numbers = _input_numbers(json.loads(paths["report"].read_text(encoding="utf-8")), BENCH)
    for slide in slides[6:11]:
        for text in _texts(slide):
            for token in NUMBER.findall(text):
                assert _traceable(token, numbers), (token, text)

    # slide 9 shows the MIA and utility values the report has
    nine = slides[8]
    shapes9 = {s.name: s for s in _shapes(nine)}
    assert "MIA" in shapes9["kpi:mia"].text_frame.text
    assert "0.51" in shapes9["kpi:mia"].text_frame.text
    assert "0.97" in shapes9["kpi:utility"].text_frame.text
    assert "utility" in shapes9["kpi:utility"].text_frame.text.lower()
    assert not {"placeholder:mia", "placeholder:utility"} & set(shapes9)

    # the steps the report lacks and the screenshots it lacks are placeholders, never figures
    assert "placeholder:screen-job" in _names(slides[6])
    assert "placeholder:screen-evidence" in _names(slides[10])
    assert "placeholder:s7" in _names(slides[7])
    assert "picture:screen-playground" in _names(slides[4])
    assert "picture:screen-scan-parallel" in _names(slides[9])
    assert result.missing  # and the build says so
    assert any("S7" in m for m in result.missing)

    # the bench file is drawn as a bar (a plain rectangle), labelled with its measured rate
    bars = [s for s in _shapes(slides[6]) if s.name.startswith("bar:")]
    assert len(bars) == 1
    assert any("941.6" in t for t in _texts(slides[6]))

    # no CRITICAL entity anywhere in the deck's words (slides and notes)
    words = [t for slide in slides for t in _texts(slide, chrome=True)]
    words += [slide.notes_slide.notes_text_frame.text for slide in slides]
    found = Detector(DetectorConfig()).detect(text="\n".join(words)).findings
    assert not [f.entity_type for f in found if f.sensitivity is Sensitivity.CRITICAL]

    assert result.out.stat().st_size < 20 * 1024 * 1024


def test_slide_nine_shows_placeholders_without_mia_and_utility(tmp_path: Path) -> None:
    facts = json.loads(json.dumps(S8_FACTS))
    facts["privacy"]["membership_inference"] = {"run": False, "reason": "holdout too small"}
    del facts["utility"]
    prs, result = _build(tmp_path, _inputs(tmp_path, {"S8": facts}))
    nine = list(prs.slides)[8]
    names = _names(nine)
    assert {"placeholder:mia", "placeholder:utility"} <= names
    assert not {"kpi:mia", "kpi:utility"} & names
    assert "kpi:fidelity" in names and "kpi:exact-rows" in names
    assert any("MIA" in m for m in result.missing)


def test_complete_report_fills_every_slot(tmp_path: Path) -> None:
    paths = _inputs(tmp_path, FULL_FACTS, screens=SCREENS_ALL)
    prs, result = _build(tmp_path, paths)
    slides = list(prs.slides)
    assert result.missing == ()
    for slide in slides:
        assert not [n for n in _names(slide) if n.startswith("placeholder:")], slide.slide_id
    numbers = _input_numbers(json.loads(paths["report"].read_text(encoding="utf-8")), BENCH)
    for slide in slides[4:11]:  # slides 5-11 carry the rehearsal's figures
        for text in _texts(slide):
            for token in NUMBER.findall(text):
                assert _traceable(token, numbers), (token, text)
    text10 = "\n".join(_texts(slides[9]))
    assert "Non-GET requests received by the site" in text10
    eleven = "\n".join(_texts(slides[10]))
    assert "5 of 11" in eleven and "raw PII sent to an LLM" in eleven
    eight = "\n".join(_texts(slides[7]))
    assert "500" in eight and "Oracle" in eight and "MongoDB" in eight and "PostgreSQL" in eight


def test_without_inputs_every_figure_is_a_placeholder(tmp_path: Path) -> None:
    out = tmp_path / "deck.pptx"
    result = build_deck.build_deck(
        out=out,
        report=tmp_path / "absent.json",
        screens=tmp_path / "absent-screens",
        bench=tmp_path / "absent-bench",
        script=SCRIPT,
    )
    slides = list(Presentation(str(out)).slides)
    assert len(slides) == 12
    for slide in slides[4:11]:
        assert [n for n in _names(slide) if n.startswith("placeholder:")]
        for text in _texts(slide):
            assert not NUMBER.findall(text), text
    assert any("report" in m for m in result.missing)
    # --strict turns any missing input into a failed build
    code = build_deck.main(
        [
            "--report",
            str(tmp_path / "absent.json"),
            "--screens",
            str(tmp_path / "absent-screens"),
            "--bench",
            str(tmp_path / "absent-bench"),
            "--script",
            str(SCRIPT),
            "--out",
            str(tmp_path / "strict.pptx"),
            "--strict",
        ]
    )
    assert code == 1


def test_a_step_that_did_not_pass_shows_no_figures(tmp_path: Path) -> None:
    paths = _inputs(tmp_path, {"S8": S8_FACTS})
    doc = json.loads(paths["report"].read_text(encoding="utf-8"))
    doc["steps"][0]["status"] = "FAIL"
    paths["report"].write_text(json.dumps(doc), encoding="utf-8")
    prs, result = _build(tmp_path, paths)
    nine = list(prs.slides)[8]
    assert not [n for n in _names(nine) if n.startswith("kpi:")]
    assert "placeholder:s8" in _names(nine)
    assert any("S8" in m and "FAIL" in m for m in result.missing)
    for text in _texts(nine):
        assert "0.93" not in text


# -- slide 1: "ran end to end" only for a rehearsal that did (S6a review r1, B1) -------------------
def _with_statuses(statuses: dict[str, Any], *, drop: tuple[str, ...] = ()) -> dict[str, Any]:
    """The complete report (S1-S11), each step PASS unless ``statuses`` says otherwise."""
    doc = _report({sid: facts for sid, facts in FULL_FACTS.items() if sid not in drop})
    for step in doc["steps"]:
        step["status"] = statuses.get(step["id"], step["status"])
    return doc


def _build_doc(base: Path, doc: dict[str, Any]) -> tuple[Any, build_deck.BuildResult]:
    """Build from ``doc`` with every screenshot and a bench file, in its own directory."""
    base.mkdir(parents=True, exist_ok=True)
    paths = _inputs(base, {}, screens=SCREENS_ALL)
    paths["report"].write_text(json.dumps(doc), encoding="utf-8")
    return _build(base, paths)


def _strict_exit(base: Path) -> int:
    """``--strict`` over the inputs :func:`_build_doc` wrote into ``base``."""
    return build_deck.main(
        [
            "--report",
            str(base / "e2e-demo-report.json"),
            "--screens",
            str(base / "demo-screens"),
            "--bench",
            str(base / "bench"),
            "--script",
            str(SCRIPT),
            "--out",
            str(base / "strict.pptx"),
            "--strict",
        ]
    )


def _title_texts(prs: Any) -> list[str]:
    return _texts(list(prs.slides)[0])


def _rehearsal_placeholder(prs: Any) -> str:
    shapes = {s.name: s for s in _shapes(list(prs.slides)[0])}
    assert "placeholder:rehearsal" in shapes
    return str(shapes["placeholder:rehearsal"].text_frame.text)


def test_title_slide_states_the_measured_end_to_end_time(tmp_path: Path) -> None:
    prs, result = _build_doc(tmp_path, _report(FULL_FACTS))
    # total_seconds 431.6 prints as 7:12 (whole seconds); the date is generated_at's
    assert (
        "Built from the rehearsal of 14 October 2026: the scripted demo ran end to end in 7:12."
        in _title_texts(prs)
    )
    assert "placeholder:rehearsal" not in _names(list(prs.slides)[0])
    assert result.missing == ()
    assert _strict_exit(tmp_path) == 0


@pytest.mark.parametrize(
    ("statuses", "reason"),
    [
        # the e2e tool stops at the first failed step and reports the later ones NOT RUN
        ({"S5": "FAIL", **{f"S{n}": "NOT RUN" for n in range(6, 12)}}, "first failure: S5"),
        ({"S5": "ERROR"}, "first failure: S5"),
        ({"S5": "BROKEN"}, "first failure: S5"),  # a status the deck does not know
        ({"S5": None}, "first failure: S5"),
        ({"S11": ["PASS"]}, "first failure: S11"),  # not even a string: unknown, never a crash
        ({"S11": "NOT RUN"}, "S11 not run"),
    ],
)
def test_title_slide_claims_no_end_to_end_run_when_a_step_did_not_pass(
    tmp_path: Path, statuses: dict[str, Any], reason: str
) -> None:
    doc = _with_statuses(statuses)
    doc["total_seconds"] = 130.2  # the wall time up to the failure is still written
    prs, result = _build_doc(tmp_path, doc)
    assert not [t for t in _title_texts(prs) if "end to end" in t or "2:10" in t]
    words = _rehearsal_placeholder(prs)
    assert reason in words and "did not complete" in words
    assert any(m.startswith("slide 1:") and reason in m for m in result.missing)
    assert _strict_exit(tmp_path) == 1


def test_title_slide_names_the_steps_the_rehearsal_skipped(tmp_path: Path) -> None:
    prs, _ = _build_doc(tmp_path / "one", _with_statuses({"S3": "SKIP"}))
    assert (
        "Built from the rehearsal of 14 October 2026: the scripted demo ran end to end in 7:12 "
        "(S3 skipped)." in _title_texts(prs)
    )
    doc = _report({**FULL_FACTS, "S12": {}})  # the console walk, listed when it ran
    for step in doc["steps"]:
        if step["id"] in {"S3", "S12"}:
            step["status"] = "SKIP"
    prs, _ = _build_doc(tmp_path / "two", doc)
    title = "\n".join(_title_texts(prs))
    assert "the scripted demo ran end to end in 7:12 (S3 and S12 skipped)." in title


def test_title_slide_needs_every_scripted_step_in_the_report(tmp_path: Path) -> None:
    prs, result = _build_doc(tmp_path / "gap", _with_statuses({}, drop=("S4",)))
    assert not [t for t in _title_texts(prs) if "end to end" in t]
    assert "no step S4 in the report" in _rehearsal_placeholder(prs)
    assert any(m.startswith("slide 1:") and "S4" in m for m in result.missing)
    doc = _report(FULL_FACTS)
    doc["steps"].append({"id": "Jane Q Customer", "status": "PASS"})  # not a step id: unreadable
    prs, _ = _build_doc(tmp_path / "odd", doc)
    assert not [t for t in _title_texts(prs) if "end to end" in t]
    assert "Jane Q Customer" not in _rehearsal_placeholder(prs)


# -- slide 11: the posture count in its context (S6a review r1, B2) -------------------------------
def _posture_kpi(prs: Any) -> tuple[str, str, str]:
    """(value text, value colour, caption) of slide 11's posture card."""
    card = {s.name: s for s in _shapes(list(prs.slides)[10])}["kpi:posture"]
    paras = card.text_frame.paragraphs
    value = paras[1].runs[0]
    caption = paras[2].runs[0].text if len(paras) > 2 and paras[2].runs else ""
    return value.text, str(value.font.color.rgb), caption


TAIL = "staging and production refuse to boot when a check enforced at boot fails"


@pytest.mark.parametrize(
    ("posture", "value", "tone", "caption"),
    [
        # the local demo stack: evaluated as if production, reported only: never the alarm colour
        (
            {"passed": 5, "total": 11, "enforced": False, "env": "local"},
            "5 of 11 passed",
            build_deck.AMBER,
            f"local stack, evaluated as if production; {TAIL}",
        ),
        (
            {"passed": 11, "total": 11, "enforced": False, "env": "dev"},
            "11 of 11 passed",
            build_deck.EMERALD,
            f"dev stack, evaluated as if production; {TAIL}",
        ),
        # no (or no known) environment name: the context still shows, no name is printed
        (
            {"passed": 5, "total": 11, "enforced": False},
            "5 of 11 passed",
            build_deck.AMBER,
            f"evaluated as if production, reported only here; {TAIL}",
        ),
        (
            {"passed": 5, "total": 11, "enforced": False, "env": "prod"},
            "5 of 11 passed",
            build_deck.AMBER,
            f"evaluated as if production, reported only here; {TAIL}",
        ),
        # an enforced failure keeps the alarm colour
        (
            {"passed": 10, "total": 11, "enforced": True, "env": "prod"},
            "10 of 11 passed",
            build_deck.ROSE,
            "enforced on this production stack: it refuses to boot when a check enforced at boot "
            "fails",
        ),
        (
            {"passed": 11, "total": 11, "enforced": True, "env": "staging"},
            "11 of 11 passed",
            build_deck.EMERALD,
            "enforced on this staging stack: it refuses to boot when a check enforced at boot fails",
        ),
    ],
)
def test_posture_count_is_shown_with_where_it_is_enforced(
    tmp_path: Path, posture: dict[str, Any], value: str, tone: str, caption: str
) -> None:
    s11 = dict(FULL_FACTS["S11"], posture=posture)
    prs, result = _build_doc(tmp_path, _report({**FULL_FACTS, "S11": s11}))
    assert _posture_kpi(prs) == (value, tone, caption)
    assert result.missing == ()
    lede = {s.name: s for s in _shapes(list(prs.slides)[10])}["lede"].text_frame.text
    assert "a posture check enforced at boot fails" in lede


@pytest.mark.parametrize("enforced", [None, "false", 0])
def test_posture_count_without_its_enforcement_is_a_placeholder(
    tmp_path: Path, enforced: Any
) -> None:
    posture: dict[str, Any] = {"passed": 5, "total": 11, "env": "local"}
    if enforced is not None:
        posture["enforced"] = enforced
    s11 = dict(FULL_FACTS["S11"], posture=posture)
    prs, result = _build_doc(tmp_path, _report({**FULL_FACTS, "S11": s11}))
    names = _names(list(prs.slides)[10])
    assert "placeholder:posture" in names and "kpi:posture" not in names
    assert any("S11 facts.posture.enforced" in m for m in result.missing)


# -- slide 10: the site's zero only with the smoke's D4 positive controls (Task 24 S4b) ------------
D4_LABEL = "Non-GET requests received by the site"
D4_CAPTION = (
    "counted by the demo site, which served the scan 175 GETs; the browser blocked 60 sends"
)


def _s9(
    *, scan: dict[str, Any] | None = None, received: dict[str, Any] | None = None, **top: Any
) -> dict[str, Any]:
    """S9's facts with ``scan`` / ``received`` keys and top-level facts replaced."""
    facts: dict[str, Any] = json.loads(json.dumps(S9_FACTS))
    facts["scan"].update(scan or {})
    facts["received"].update(received or {})
    facts.update(top)
    return facts


def _scanner_kpi(prs: Any) -> tuple[str, str, str, str]:
    """(label, value text, value colour, caption) of slide 10's site-counter card."""
    card = {s.name: s for s in _shapes(list(prs.slides)[9])}["kpi:posts"]
    paras = card.text_frame.paragraphs
    value = paras[1].runs[0]
    caption = paras[2].runs[0].text if len(paras) > 2 and paras[2].runs else ""
    return paras[0].runs[0].text, value.text, str(value.font.color.rgb), caption


def test_scanner_zero_is_shown_with_its_positive_controls(tmp_path: Path) -> None:
    prs, result = _build_doc(tmp_path, _report(FULL_FACTS))
    assert _scanner_kpi(prs) == (D4_LABEL, "0", build_deck.EMERALD, D4_CAPTION)
    assert result.missing == ()
    assert _strict_exit(tmp_path) == 0


@pytest.mark.parametrize(
    ("facts", "why"),
    [
        # the site never served the scan: these counters are not the scanned site's
        (_s9(received={"get": 0}), "S9 D4 positive controls"),
        # fewer GETs than pages visited (every visit fetches its page: it sends no-store)
        (_s9(received={"get": 20}), "S9 D4 positive controls"),
        # the guard stopped no send: the scan never provoked one, so the zero proves nothing
        (_s9(sends_blocked=0), "S9 D4 positive controls"),
        # the contract before the scanner merged: a bare count, no positive controls
        (
            {"scan": dict(S9_FACTS["scan"]), "posts_received": 0},
            "S9 facts.received.non_get",
        ),
        (_s9(received={"form_get": None}), "S9 facts.received.form_get"),
        (_s9(scan={"non_get_allowed": "0"}), "S9 facts.scan.non_get_allowed"),
        ({k: v for k, v in _s9().items() if k != "sends_blocked"}, "S9 facts.sends_blocked"),
    ],
)
def test_scanner_zero_without_its_proof_is_a_placeholder(
    tmp_path: Path, facts: dict[str, Any], why: str
) -> None:
    prs, result = _build_doc(tmp_path, _report({**FULL_FACTS, "S9": facts}))
    names = _names(list(prs.slides)[9])
    assert "placeholder:posts" in names and "kpi:posts" not in names
    assert any(why in m for m in result.missing), result.missing
    assert _strict_exit(tmp_path) == 1


@pytest.mark.parametrize(
    ("facts", "value"),
    [
        (_s9(received={"non_get": 2}), "2"),
        (_s9(received={"form_get": 1}), "0"),  # a form sent by GET is a submission too
        (_s9(received={"admin_hits": 1}), "0"),  # a denied path reached the site
        (_s9(scan={"non_get_allowed": 1}), "0"),  # the request log let one through
    ],
)
def test_scanner_count_that_moved_is_never_emerald(
    tmp_path: Path, facts: dict[str, Any], value: str
) -> None:
    prs, _ = _build_doc(tmp_path, _report({**FULL_FACTS, "S9": facts}))
    label, shown, tone, caption = _scanner_kpi(prs)
    assert (label, shown, tone) == (D4_LABEL, value, build_deck.ROSE)
    assert caption == "the demo site counted a request that could submit: S9 must fail"


def test_strings_that_are_not_names_never_reach_a_slide(tmp_path: Path) -> None:
    """The report is value-free by contract; the deck still prints only name-shaped strings."""
    s2 = dict(FULL_FACTS["S2"], entity_types=["EMAIL", "123-45-6789", "Jane Q Customer"])
    paths = _inputs(
        tmp_path,
        {"S2": s2},
        report_top={"tenant": "Jane Q Customer 123-45-6789", "commit": "not a sha!"},
    )
    prs, _ = _build(tmp_path, paths)
    words = "\n".join(t for slide in prs.slides for t in _texts(slide, chrome=True))
    assert "123-45-6789" not in words and "Jane Q Customer" not in words
    assert "not a sha" not in words
    assert "EMAIL" in words


def test_facts_of_an_unexpected_shape_become_placeholders(tmp_path: Path) -> None:
    steps = {
        "S2": dict(FULL_FACTS["S2"], entity_types=["not a name"], findings="13"),
        "S6": {"jobs": [{"engine": "sqlite", "counts": {"rows_written": 5}}]},
        "S7": dict(FULL_FACTS["S7"], pairs=[{"left": 0, "right": 9, "checked": 1, "matched": 1}]),
    }
    prs, result = _build(tmp_path, _inputs(tmp_path, steps, screens=SCREENS_ALL))
    slides = list(prs.slides)
    assert {"placeholder:findings", "placeholder:entity-types"} <= _names(slides[4])
    assert "placeholder:s6" in _names(slides[6])
    assert "placeholder:s7-pairs" in _names(slides[7])
    assert not [n for n in _names(slides[7]) if n.startswith("pair:")]
    assert any("S2 facts.findings" in m for m in result.missing)


def test_a_report_of_another_shape_is_refused(tmp_path: Path) -> None:
    bad = tmp_path / "report.json"
    bad.write_text(json.dumps({"schema": "something-else", "steps": []}), encoding="utf-8")
    with pytest.raises(build_deck.DeckInputError):
        build_deck.build_deck(out=tmp_path / "x.pptx", report=bad, script=SCRIPT)
    bad.write_text("{not json", encoding="utf-8")
    code = build_deck.main(["--report", str(bad), "--out", str(tmp_path / "y.pptx")])
    assert code == 2
    assert not (tmp_path / "y.pptx").exists()


def test_speaker_notes_come_from_the_demo_script(tmp_path: Path) -> None:
    notes = build_deck.script_notes(SCRIPT)
    assert sorted(notes) == list(range(1, 13))
    assert all(text.strip() for text in notes.values())
    prs, result = _build(tmp_path, _inputs(tmp_path, {"S8": S8_FACTS}))
    for number, slide in enumerate(prs.slides, start=1):
        assert slide.notes_slide.notes_text_frame.text == notes[number]
    assert not any("script" in m for m in result.missing)


def test_script_notes_join_wrapped_lines_and_skip_code(tmp_path: Path) -> None:
    script = tmp_path / "script.md"
    script.write_text(
        "# Title\n\n"
        "#### Slide 2 · Two\n\n"
        '> "First spoken line that wraps\n'
        '> onto a second line."\n'
        '> "Second **line** with `code`."\n'
        ">\n"
        "> - a list item\n\n"
        "```bash\n"
        "# a shell comment inside a fence is not a heading\n"
        "echo nothing\n"
        "```\n"
        "Still slide two.\n\n"
        "## Next beat\n\n"
        "Not a note.\n\n"
        "#### Slide 13 · out of range\n\nignored\n",
        encoding="utf-8",
    )
    assert build_deck.script_notes(script) == {
        2: "First spoken line that wraps onto a second line.\n"
        "Second line with code.\n"
        "• a list item\n"
        "Still slide two."
    }


def test_without_the_script_the_notes_fall_back_and_say_so(tmp_path: Path) -> None:
    out = tmp_path / "deck.pptx"
    result = build_deck.build_deck(out=out, script=tmp_path / "no-script.md")
    for slide in Presentation(str(out)).slides:
        assert slide.notes_slide.notes_text_frame.text.strip()
    assert any("script" in m for m in result.missing)


def test_cli_builds_the_deck(tmp_path: Path) -> None:
    paths = _inputs(tmp_path, {"S8": S8_FACTS, "S9": S9_FACTS})
    out = tmp_path / "cli.pptx"
    code = build_deck.main(
        [
            "--report",
            str(paths["report"]),
            "--screens",
            str(paths["screens"]),
            "--bench",
            str(paths["bench"]),
            "--script",
            str(SCRIPT),
            "--out",
            str(out),
        ]
    )
    assert code == 0 and out.stat().st_size > 0


# -- the committed deck (Task 26 S6b: built from a complete rehearsal) ----------------------------
COMMITTED = ROOT / "docs" / "demo" / "dps-showcase.pptx"


def test_the_committed_deck_is_a_complete_keynote_safe_build() -> None:
    """``docs/demo/dps-showcase.pptx`` as committed: twelve slides with titles and the script's
    notes, no placeholder (it was built ``--strict`` from a complete rehearsal), no chart part or
    graphic frame, only Helvetica Neue and Arial, no CRITICAL entity in its words, under 20 MB.
    Rebuild it (``python tools/build_deck.py --strict``) after editing a Slide block."""
    if not COMMITTED.is_file():
        pytest.fail("docs/demo/dps-showcase.pptx is missing: rehearse, then build it --strict")
    prs = Presentation(str(COMMITTED))
    slides = list(prs.slides)
    assert len(slides) == 12
    notes = build_deck.script_notes(SCRIPT)
    for number, slide in enumerate(slides, start=1):
        assert slide.shapes.title is not None and slide.shapes.title.text.strip()
        assert slide.notes_slide.notes_text_frame.text == notes[number], f"slide {number} notes"
        assert not [n for n in _names(slide) if n.startswith("placeholder:")], number
    with zipfile.ZipFile(COMMITTED) as z:
        names = z.namelist()
        assert not [n for n in names if n.startswith(("ppt/charts/", "ppt/embeddings/"))]
        xml = {n: z.read(n).decode("utf-8") for n in names if n.endswith(".xml")}
    for name, text in xml.items():
        if name.startswith("ppt/slides/slide"):
            assert "<p:graphicFrame" not in text, name
        faces = set(re.findall(r'typeface="([^"]*)"', text))
        assert faces <= ALLOWED_FONTS | {""} | {f for f in faces if f.startswith("+m")}, name
    words = [t for slide in slides for t in _texts(slide, chrome=True)]
    words += [slide.notes_slide.notes_text_frame.text for slide in slides]
    found = Detector(DetectorConfig()).detect(text="\n".join(words)).findings
    assert not [f.entity_type for f in found if f.sensitivity is Sensitivity.CRITICAL]
    assert COMMITTED.stat().st_size < 20 * 1024 * 1024

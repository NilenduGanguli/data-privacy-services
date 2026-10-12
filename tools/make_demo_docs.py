"""Write the deterministic synthetic demo documents (no real data) for the console, the smoke test
and the ten-minute demo.

    python tools/make_demo_docs.py            # → demo/documents/*, frontend/public/demo/kyc-scanned.pdf

Every value comes from the surrogate engine (seed 7): personas, a Luhn-valid card on the test IIN,
RFC 2606 e-mail domains. The Aadhaar-shaped number is Verhoeff arithmetic from the same seed on an
issued first digit (2-9), as in the evaluation corpus: the surrogate engine's start with 1 (never
issued), which the detector rightly ignores, so the demo letter would show it unboxed. Re-running
produces identical files, so they are committed (``tests/test_demo_docs.py`` keeps them current).
"""

from __future__ import annotations

import csv
import hashlib
import io
import random
import shutil
import zipfile
from pathlib import Path

import docx
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from dps.engine.detect.checksums import verhoeff_check_digit
from dps.engine.synth.surrogates import SurrogateEngine
from dps.layers.documents.redact import build_image_pdf
from dps.layers.documents.render import PdfRenderer

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "demo" / "documents"
PUBLIC = ROOT / "frontend" / "public" / "demo"
SEED = hashlib.sha256(b"dps.demo.documents|7").digest()


def _aadhaar(seed: bytes) -> str:
    """A Verhoeff-valid Aadhaar-shaped number on an issued first digit (2-9), from ``seed``."""
    rng = random.Random(int.from_bytes(hashlib.sha256(seed + b"|aadhaar").digest(), "big"))
    body = str(rng.randint(2, 9)) + "".join(str(rng.randint(0, 9)) for _ in range(10))
    number = body + verhoeff_check_digit(body)
    return f"{number[:4]} {number[4:8]} {number[8:]}"


def _values() -> dict[str, str]:
    engine = SurrogateEngine(SEED, default_locale="en_IN")
    persona = engine.persona("kyc", "en_IN")
    return {
        "name": persona.full_name,
        "street": persona.street,
        "city": persona.city,
        "pin": persona.postal_code,
        "dob": persona.dob.strftime("%d/%m/%Y"),
        "gender": persona.gender.title(),
        "email": engine.surrogate("EMAIL", "x", "en_IN", "kyc"),
        "phone": engine.surrogate("PHONE", "+91 98451 20377", "en_IN", "kyc"),
        "aadhaar": _aadhaar(SEED),
        "pan": engine.surrogate("PAN_IN", "AAAAA0000A", "en_IN", "kyc"),
        "card": engine.surrogate("CARD_PAN", "0000 0000 0000 0000", "en_IN", "kyc"),
        "account": engine.surrogate("ACCOUNT_NUMBER", "0000 0000 0000 00", "en_IN", "kyc"),
        "payee": engine.persona("payee", "en_IN").full_name,
    }


def kyc_pdf(v: dict[str, str]) -> bytes:
    """Two A4 pages: an Aadhaar-style letter and a savings statement (text layer)."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    width, height = A4
    pages = [
        (
            "Unique Identification Letter",
            [
                f"To: {v['name']}",
                f"Address: {v['street']}, {v['city']} PIN code {v['pin']}",
                f"Mobile: {v['phone']}",
                f"Your Aadhaar No.: {v['aadhaar']}",
                f"DOB: {v['dob']}   Gender: {v['gender']}",
                "Aadhaar is proof of identity, not of citizenship.",
            ],
        ),
        (
            "Savings account statement",
            [
                f"Customer: {v['name']}",
                f"Email: {v['email']}   Phone: {v['phone']}",
                f"Account no.: {v['account']}   Customer ID: CUST-884120",
                f"PAN: {v['pan']}",
                f"02 Sep  NEFT/{v['payee']}/SALARY ADV            45,000.00",
                f"07 Sep  POS card {v['card']} BLR      6,120.00",
                "30 Sep  Closing balance                       1,16,946.50",
            ],
        ),
    ]
    for title, lines in pages:
        c.setFont("Helvetica-Bold", 18)
        c.drawString(56, height - 72, title)
        c.setFont("Helvetica", 12)
        for i, line in enumerate(lines):
            c.drawString(56, height - 120 - i * 26, line)
        c.showPage()
    c.save()
    return buf.getvalue()


def scanned(pdf: bytes, *, dpi: int = 200) -> bytes:
    """The same pages as JPEG rasters only — a 'scanned' PDF with no text layer."""
    with PdfRenderer(pdf) as renderer:
        pages = [(renderer.render(i, dpi), *renderer.page_size(i)) for i in range(len(renderer))]
    return build_image_pdf(pages, image_format="jpeg", jpeg_quality=80)


def id_png(pdf: bytes, *, dpi: int = 200) -> bytes:
    """Page 1 as a PNG (an image upload: OCR reads it)."""
    with PdfRenderer(pdf) as renderer:
        image = renderer.render(0, dpi)
    buf = io.BytesIO()
    image.save(buf, "PNG", dpi=(dpi, dpi))
    return buf.getvalue()


def kyc_docx(v: dict[str, str]) -> bytes:
    """A KYC review note: a header, a paragraph and a two-column table."""
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "KYC review · Customer ID CUST-884120"
    document.add_paragraph(
        f"Customer {v['name']} ({v['email']}) asked to update the PAN {v['pan']}."
    )
    table = document.add_table(rows=3, cols=2)
    for row, (label, value) in enumerate(
        (("Aadhaar", v["aadhaar"]), ("Card", v["card"]), ("Phone", v["phone"]))
    ):
        table.cell(row, 0).text, table.cell(row, 1).text = label, value
    buf = io.BytesIO()
    document.save(buf)
    return _stable_zip(buf.getvalue())


def _stable_zip(data: bytes) -> bytes:
    """Re-pack a zip with fixed timestamps so the committed .docx is byte-for-byte reproducible."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(fixed, source.read(info.filename))
    return out.getvalue()


def customers_csv(rows: int = 20) -> bytes:
    """``rows`` customers, one persona each."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(
        ["customer_id", "full_name", "email", "phone", "card_number", "status", "risk_score"]
    )
    for i in range(rows):
        engine = SurrogateEngine(
            hashlib.sha256(SEED + f"|row{i}".encode()).digest(), default_locale="en_IN"
        )
        persona = engine.persona("row", "en_IN")
        writer.writerow(
            [
                f"CUST-{884120 + i}",
                persona.full_name,
                engine.surrogate("EMAIL", "x", "en_IN", "row"),
                engine.surrogate("PHONE", "+91 98451 20377", "en_IN", "row"),
                engine.surrogate("CARD_PAN", "0000 0000 0000 0000", "en_IN", "row"),
                "ACTIVE" if i % 4 else "CLOSED",
                600 + (i * 37) % 300,
            ]
        )
    return buf.getvalue().encode("utf-8")


def build() -> dict[str, bytes]:
    """Every demo document by file name (nothing is written)."""
    v = _values()
    text = kyc_pdf(v)
    return {
        "kyc-text.pdf": text,
        "kyc-scanned.pdf": scanned(text),
        "id-card.png": id_png(text),
        "kyc.docx": kyc_docx(v),
        "customers.csv": customers_csv(),
        "kyc.txt": (
            f"Customer {v['name']} (DOB: {v['dob']}) mailed {v['email']} about card {v['card']}.\n"
            f"Aadhaar {v['aadhaar']} and PAN {v['pan']} were verified.\n"
        ).encode(),
    }


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    PUBLIC.mkdir(parents=True, exist_ok=True)
    for name, data in build().items():
        (OUT / name).write_bytes(data)
        print(f"wrote demo/documents/{name} ({len(data)} bytes)")
    shutil.copyfile(OUT / "kyc-scanned.pdf", PUBLIC / "kyc-scanned.pdf")
    print("wrote frontend/public/demo/kyc-scanned.pdf")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
